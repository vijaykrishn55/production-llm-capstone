from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown"}


def parse_document(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}. Use PDF, DOCX, or Markdown.")
    if suffix == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif suffix == ".docx":
        document = DocxDocument(str(path))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        tables = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(paragraphs + tables)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    text = text.strip()
    if not text:
        raise ValueError(f"No extractable text found in {path.name}; scanned PDFs need OCR first.")
    return text
